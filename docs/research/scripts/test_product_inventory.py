"""Consistency checks for the research package; run with Python's unittest.

python3 -m unittest discover -s docs/research/scripts -p 'test_*.py'
The optional workbook checks require openpyxl; CSV/Markdown checks do not.
"""
import collections
import csv
from pathlib import Path
import re
import unittest

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "docs/research/pdf-product-strategy-2026-10-25"


def rows(name):
    with (OUT / name).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


class ResearchPackageTests(unittest.TestCase):
    def test_complete_baseline_and_gap_subset(self):
        baseline = rows("baseline-features.csv")
        gaps = rows("existing-gaps.csv")
        self.assertEqual(len(baseline), 833)
        self.assertEqual(len(gaps), 384)
        self.assertEqual(collections.Counter(r["repo_status"] for r in baseline), {
            "shipped": 426, "partial": 76, "planned": 308, "na": 23,
        })
        self.assertEqual(
            {r["feature_id"] for r in gaps},
            {r["feature_id"] for r in baseline if r["repo_status"] in {"planned", "partial"}},
        )
        self.assertTrue(all(r["evidence_paths"].startswith("all referenced") for r in baseline))
        for row in baseline:
            self.assertNotIn(None, row)
            self.assertNotIn(None, row.values())

    def test_recommendation_references_and_category_counts(self):
        recommendations = rows("recommended-roadmap.csv")
        source_ids = {r["source_id"] for r in rows("sources.csv")}
        feature_ids = {r["feature_id"] for r in rows("baseline-features.csv")}
        self.assertEqual(len(recommendations), 222)
        self.assertEqual(len({r["recommendation_id"] for r in recommendations}), 222)
        self.assertEqual(len(source_ids), 29)
        self.assertEqual(len({r["category"] for r in recommendations}), 20)
        self.assertEqual(collections.Counter(r["priority"] for r in recommendations), {
            "P0": 38, "P1": 98, "P2": 71, "P3": 15,
        })
        for row in recommendations:
            self.assertNotIn(None, row)
            self.assertTrue(set(row["source_ids"].split()) <= source_ids)
            ids = row["baseline_ids"]
            if ids != "'-":  # CSV protects a standalone minus from formula interpretation.
                self.assertTrue(set(ids.split()) <= feature_ids, row["recommendation_id"])
            self.assertTrue(row["done_when"])
            self.assertTrue(row["dependencies"])

    def test_all_local_markdown_links_resolve(self):
        for path in (ROOT / "docs/research").rglob("*.md"):
            for target in re.findall(r"\]\(([^)]+)\)", path.read_text()):
                if target.startswith(("https://", "http://", "#")):
                    continue
                self.assertTrue((path.parent / target.split("#")[0]).exists(), (path, target))

    def test_markdown_has_every_recommendation_and_gap(self):
        recommendations = (OUT / "recommended-roadmap.md").read_text()
        gaps = (OUT / "existing-gaps.md").read_text()
        for row in rows("recommended-roadmap.csv"):
            self.assertEqual(recommendations.count("| " + row["recommendation_id"] + " |"), 1)
        for row in rows("existing-gaps.csv"):
            self.assertEqual(gaps.count("| " + row["feature_id"] + " |"), 1)

    def test_workbook_dimensions_filters_and_text_cells(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("Install openpyxl to validate the optional workbook")
        workbook = load_workbook(OUT / "Linkco-PDF-Product-Strategy.xlsx")
        self.assertEqual(workbook.sheetnames, ["Read me", "Recommendations", "Baseline 833", "Gaps 384", "Sources"])
        for sheet, count in [("Recommendations", 222), ("Baseline 833", 833), ("Gaps 384", 384), ("Sources", 29)]:
            ws = workbook[sheet]
            self.assertEqual(ws.max_row, count + 1)
            self.assertEqual(ws.freeze_panes, "B2")
            self.assertEqual(ws.auto_filter.ref, ws.dimensions)
            self.assertTrue(all(cell.data_type != "f" for row in ws for cell in row))
        workbook.close()


if __name__ == "__main__":
    unittest.main()
