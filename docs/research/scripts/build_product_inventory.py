#!/usr/bin/env python3
"""Rebuild the dated PDF market-study inventory from its pinned source snapshot.

Standard library generates CSV/Markdown; --xlsx additionally needs openpyxl.
Usage: python3 docs/research/scripts/build_product_inventory.py [--xlsx]
This is documentation tooling, not part of the product or Rust build.
"""
from __future__ import annotations

import argparse
import collections
import csv
from pathlib import Path
import re
import subprocess
import tomllib
import zipfile

ROOT = Path(__file__).resolve().parents[3]
BASELINE = "1d7f56366160011e5bf9e370646bf1b45ba5fc79"
OUT = ROOT / "docs/research/pdf-product-strategy-2026-10-25"
INPUT = Path(__file__).with_name("recommendations_2026_10_25.psv")


def git_text(path: str) -> str:
    return subprocess.check_output(["git", "show", f"{BASELINE}:{path}"], cwd=ROOT).decode()


def safe_csv(value: object) -> str:
    text = str(value)
    # Avoid spreadsheet formula interpretation when opening CSV. The workbook uses text cells.
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export_csv(name: str, rows: list[dict], fields: list[str]) -> None:
    with (OUT / name).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        for row in rows:
            w.writerow({key: safe_csv(row.get(key, "")) for key in fields})


def md(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", action="store_true")
    args = parser.parse_args()
    source = tomllib.loads(git_text("parity/acrobat-features.toml"))
    features = source["feature"]
    areas = source["schema"]["areas"]
    by_id = {f["id"]: f for f in features}
    assert len(by_id) == len(features) == 833
    tracked = set(subprocess.check_output(["git", "ls-tree", "-r", "--name-only", BASELINE], cwd=ROOT).decode().splitlines())
    inventory = []
    for f in features:
        absent = [e for e in f.get("evidence", []) if "/" in e.split("::")[0] and e.split("::")[0] not in tracked]
        inventory.append({
            "feature_id": f["id"], "area": areas[f["area"]], "title": f["title"],
            "repo_status": f["status"], "repo_tier": f["tier"], "milestone": f["milestone"],
            "notes": f.get("notes", ""), "commands": "; ".join(f.get("commands", [])),
            "tools": "; ".join(f.get("tools", [])), "evidence": "; ".join(f.get("evidence", [])),
            "evidence_paths": "missing: " + "; ".join(absent) if absent else "all referenced paths exist; tests not rerun in this study",
            "baseline_commit": BASELINE,
        })
    with INPUT.open(encoding="utf-8", newline="") as f:
        raw = list(csv.DictReader(f, delimiter="|"))
    sources = list(csv.DictReader((OUT / "sources.csv").open(encoding="utf-8", newline="")))
    source_ids = {s["source_id"] for s in sources}
    counters: collections.Counter = collections.Counter()
    recs = []
    for row in raw:
        assert None not in row and all(v is not None for v in row.values()), row
        prefix = row["category"].split()[0]
        counters[prefix] += 1
        ids = [] if row["baseline_ids"] == "-" else row["baseline_ids"].split()
        assert set(ids) <= by_id.keys(), (row["function"], ids)
        assert set(row["source_ids"].split()) <= source_ids
        assert set(row["surfaces"].split()) <= {"Desktop", "Web", "CLI", "Integrations", "Mobile", "Service"}
        status = "; ".join(f"{i}: {by_id[i]['status']}" for i in ids) if ids else "proposal / enhancement; no exact parity mapping; verify scope before implementation"
        recs.append({"recommendation_id": f"{prefix}{counters[prefix]:02d}", **row, "baseline_status": status})
    assert len(recs) == 222 and len(sources) == 29
    assert len({r["recommendation_id"] for r in recs}) == len(recs)
    gaps = [f for f in inventory if f["repo_status"] in {"planned", "partial"}]
    assert len(gaps) == 384
    OUT.mkdir(parents=True, exist_ok=True)
    fields = list(inventory[0])
    export_csv("baseline-features.csv", inventory, fields)
    export_csv("existing-gaps.csv", gaps, fields)
    rec_fields = ["recommendation_id", "category", "function", "priority", "surfaces", "baseline_status", "baseline_ids", "effort", "dependencies", "done_when", "source_ids"]
    export_csv("recommended-roadmap.csv", recs, rec_fields)
    counts = collections.Counter(f["repo_status"] for f in inventory)
    tiers = collections.Counter(r["priority"] for r in recs)
    text = ["# Detailed recommended roadmap\n", "> Research recommendations, not delivered features. Priorities are this study's priorities, not the original checklist tiers.\n",
            "Read [the strategy report](README.md) first for scope, definitions, caveats and sequencing. Source IDs refer to [sources.md](sources.md). A source supports the workflow rationale, not our current capability. Baseline status is the exact checklist mapping; unlinked enhancements are not asserted absent.\n",
            "Desktop = Windows/macOS/Linux/FreeBSD application; Web = browser application; CLI = command line/MCP; Integrations = OS/storage connectors; Mobile and Service = proposed product surfaces.\n",
            "P0 = non-negotiable safety/usability gate when that surface or workflow is offered; P1 = daily-work priority; P2 = advanced/segment need; P3 = future investment requiring validation.\n",
            "Effort is relative only: S = localized change; M = one bounded subsystem; L = multiple subsystems; XL = new platform/service or major fidelity work. These are not estimates in days or delivery commitments.\n",
            f"**{len(recs)} recommendations:** " + ", ".join(f"{k}: {tiers[k]}" for k in sorted(tiers)) + ".\n"]
    for category in dict.fromkeys(r["category"] for r in recs):
        text += [f"\n## {category}\n", "| ID | Function | Priority / surface / effort | Baseline mapping | Dependency | Acceptance criterion | Basis |", "|---|---|---|---|---|---|---|"]
        for r in recs:
            if r["category"] == category:
                text.append("| " + " | ".join(map(md, [r["recommendation_id"], r["function"], f"{r['priority']} · {r['surfaces']} · {r['effort']}", r["baseline_status"], r["dependencies"], r["done_when"], r["source_ids"]])) + " |")
    (OUT / "recommended-roadmap.md").write_text("\n".join(text).rstrip() + "\n", encoding="utf-8")
    text = ["# Complete baseline gap register\n", f"Snapshot: `{BASELINE}`. These **{len(gaps)}** entries comprise all **308 planned** and **76 partial** entries in the 833-item repository checklist. These are repository assertions, not fresh UI verification.\n",
            "No out-of-scope (`na`) item is silently counted as missing. The 23 `na` entries and all 426 shipped entries remain visible in [the complete baseline CSV](baseline-features.csv). See the main report for known contradictions and commercial scope decisions.\n"]
    for area in areas.values():
        sub = [f for f in gaps if f["area"] == area]
        text += [f"\n## {area} ({len(sub)})\n", "| Feature ID | Function | Repository status / tier | Exact limitation notes |", "|---|---|---|---|"]
        text.extend("| " + " | ".join(map(md, [f["feature_id"], f["title"], f"{f['repo_status']} / {f['repo_tier']}", f["notes"] or "No additional limitation recorded; requires code/UI verification."])) + " |" for f in sub)
    (OUT / "existing-gaps.md").write_text("\n".join(text).rstrip() + "\n", encoding="utf-8")
    text = ["# Sources and research boundaries\n", "Reviewed 2026-10-25. Vendor pages establish advertised workflows, not measured adoption, market share, independent product quality or Linkco demand. No customer interviews or hands-on competitor benchmark was performed. Plan-specific descriptions may change. Full pages were fetched for selected sources; other entries are official search extracts, as marked below.\n"]
    for s in sources:
        citation = f"[{s['search_result_id']}]({s['url']})"
        text += [f"\n## {s['source_id']} — {s['publisher']}: {s['title']}\n", f"{citation}\n", f"- Review method: {s['review_method']}.", f"- Supports: {s['what_it_supports']}.", f"- Limits: {s['limitations']}."]
    text += ["\n## Repository sources\n", f"Baseline commit: `{BASELINE}`; `parity/acrobat-features.toml`, README, architecture and crate/module sources listed in the report. All referenced evidence file paths exist at this snapshot. Existence does not prove the referenced function is a test or that it passed. This study did not rerun Cargo or a GUI session.\n"]
    (OUT / "sources.md").write_text("\n".join(text).rstrip() + "\n", encoding="utf-8")
    if args.xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        summary = wb.active
        summary.title = "Read me"
        summary_rows = [
            ["Linkco PDF product and market study", "2026-10-25"],
            ["Read first", "README.md is the decision report; this workbook is the filterable inventory."],
            ["Baseline", BASELINE], ["Existing baseline features", len(inventory)],
            ["Repository status counts", str(dict(counts))], ["Existing incomplete features", len(gaps)],
            ["Recommendations", len(recs)], ["Recommendation priorities", str(dict(tiers))],
            ["Not additive", "222 recommendations overlap each other and the 833 baseline entries; do not add these counts."],
            ["Status limits", "Shipped/partial/planned are repository claims; no fresh runtime verification in this study."],
            ["New surfaces", "Mobile and Service are proposals, not existing apps."],
            ["Priorities", "P0 safety/usability when offered; P1 daily; P2 advanced; P3 future."],
            ["Effort", "S/M/L/XL are relative scope, not dates or quotations."],
            ["No exact mapping", "An unlinked recommendation may enhance existing code; it is not necessarily missing."],
            ["Source IDs", "Sources support rationale; they do not certify our functionality or legal compliance."],
        ]
        for row in summary_rows:
            summary.append(row)
        for name, rows, columns in [("Recommendations", recs, rec_fields), ("Baseline 833", inventory, fields), ("Gaps 384", gaps, fields), ("Sources", sources, list(sources[0]))]:
            ws = wb.create_sheet(name)
            ws.append(columns)
            for row in rows:
                ws.append([str(row.get(c, "")) for c in columns])
            ws.auto_filter.ref = ws.dimensions
        for ws in wb:
            ws.freeze_panes = "B2"
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="17344D")
                cell.font = Font(color="FFFFFF", bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.row_dimensions[row[0].row].height = 58
            for idx, col in enumerate(ws.iter_cols(), 1):
                label = str(ws.cell(1, idx).value)
                ws.column_dimensions[get_column_letter(idx)].width = 65 if label in {"notes", "done_when", "evidence", "baseline_status", "url"} else 32
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        summary.column_dimensions["A"].width = 38
        summary.column_dimensions["B"].width = 100
        workbook_path = OUT / "Linkco-PDF-Product-Strategy.xlsx"
        wb.properties.creator = "Linkco PDF contributors"
        wb.properties.lastModifiedBy = "Linkco PDF contributors"
        wb.save(workbook_path)
        # Pin package dates so regenerating this historical snapshot is byte-stable
        # with the same openpyxl version (3.1.5 used for this study).
        with zipfile.ZipFile(workbook_path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        core = members["docProps/core.xml"].decode()
        core = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]+", r"\g<1>2026-10-25T00:00:00Z", core)
        members["docProps/core.xml"] = core.encode()
        with zipfile.ZipFile(workbook_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name in sorted(members):
                entry = zipfile.ZipInfo(name, date_time=(2026, 10, 25, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(entry, members[name])
    print(f"Validated {len(inventory)} baseline features, {len(gaps)} incomplete entries, {len(recs)} recommendations and {len(sources)} sources.")
    print("Recommendation priorities:", dict(sorted(tiers.items())))


if __name__ == "__main__":
    main()
